import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.chart import BarChart, Reference

# ==========================================
# Stage 0 : Settings (設定區)
# ==========================================
data_dir = Path(r"C:\ntutdata\debt")
sic_excel_path = Path(
    r"C:\Users\olivi\Desktop\acer\ntut\1Seminar\data\GVKEY_debt structure_with"
    r" SIC.xlsx"
)
output_analysis_excel = Path(
    r"C:\Users\olivi\Desktop\acer\ntut\1Seminar\data\Debt_Structure_Annual_Analysis.xlsx"
)
output_debug_excel = Path(
    r"C:\Users\olivi\Desktop\acer\ntut\1Seminar\data\Debt_Debug_Report.xlsx"
)

# 彈性 Debug 年度設定：若設為 None 則稽核全年度；若設為 2008 則 Debug 專注於 2008
DEBUG_YEAR = None

print("=" * 80)
print("Debt Structure Annual Analysis + Comprehensive Debug Tool")
print("=" * 80)


# ==========================================
# Helper Functions (工具函數區)
# ==========================================
def extract_column_name(col):
    match = re.match(r"\((.*?)\)", str(col))
    if match:
        return match.group(1).strip().lower()
    return str(col).strip().lower()


def clean_to_string(val):
    if pd.isna(val):
        return None
    txt = str(val).strip()
    if txt.lower() in ["", "nan", "null"]:
        return None
    return txt.split(".")[0]


def parse_year(date_val):
    if pd.isna(date_val):
        return None
    txt = str(date_val).strip()
    m = re.search(r"(19\d{2}|20\d{2})", txt)
    if m:
        return int(m.group())
    try:
        d = pd.to_datetime(txt, errors="coerce")
        if pd.notna(d):
            return d.year
    except:
        pass
    return None


def calc_actual_value(value, unit):
    try:
        value = float(value)
    except:
        value = 0.0
    try:
        unit = int(unit)
    except:
        unit = 0
    if unit == 1:
        return value * 1000
    elif unit == 2:
        return value * 1000000
    return value


def map_debt_category(subtype, level):
    try:
        subtype = int(subtype)
    except:
        subtype = -1
    try:
        level = int(level)
    except:
        level = -1

    if subtype == 1:
        return "CP"
    elif subtype == 2:
        return "DC"
    elif subtype == 3:
        return "TL"
    elif subtype == 5:
        return "CL"
    elif subtype in [6, 7, 9]:
        return "Others"
    elif subtype == 4:
        if level == 1:
            return "SBN"
        elif level in [2, 3, 4, 5, 6, 7]:
            return "SUB"

    return None


print("Helper Functions Loaded Successfully.")
print()

# ==========================================
# Stage 1 : Processing SIC Mapping (排除金融、公用事業及 SIC 空值)
# ==========================================
print("=" * 80)
print("Stage 1 : Processing SIC Mapping")
print("=" * 80)

try:
    sic_df = pd.read_excel(sic_excel_path)
    sic_df.columns = [extract_column_name(c) for c in sic_df.columns]
    print(f"SIC File Loaded : {len(sic_df):,} rows")

    required_columns = ["companyid", "sic"]
    for col in required_columns:
        if col not in sic_df.columns:
            raise Exception(f"Missing column : {col}")

    has_gvkey = "gvkey" in sic_df.columns
    sic_df["companyid"] = sic_df["companyid"].apply(clean_to_string)

    # 1. 嚴格同時清除 companyid 與 sic 任何一個欄位為空值的列
    before_count = len(sic_df)
    sic_df = sic_df.dropna(subset=["companyid", "sic"])
    after_blank = len(sic_df)
    print(f"Remove Blank CompanyID / SIC : {before_count - after_blank:,} rows")

    # 2. 轉為數值型態
    sic_df["sic"] = pd.to_numeric(sic_df["sic"], errors="coerce")
    sic_df = sic_df.dropna(subset=["sic"])
    sic_df["sic"] = sic_df["sic"].astype(int)

    # 3. 嚴格排除公用事業 (4900-4949) 與 金融業 (6000-6999)
    filtered_sic_df = sic_df[
        ~(
            ((sic_df["sic"] >= 4900) & (sic_df["sic"] <= 4949))
            | ((sic_df["sic"] >= 6000) & (sic_df["sic"] <= 6999))
        )
    ].copy()
    print(f"Company After SIC Filter : {len(filtered_sic_df):,}")

    valid_companies = set(filtered_sic_df["companyid"].unique())
    print(f"Whitelist Unique CompanyIDs : {len(valid_companies):,}")

    # 4. 建立 CompanyID → GVKEY 映射字典（僅限定在有效白名單公司）
    company_to_gvkey = {}
    if has_gvkey:
        filtered_sic_df["gvkey"] = filtered_sic_df["gvkey"].apply(clean_to_string)
        # 過濾掉 gvkey 為空的項目建立 dict
        valid_gvkey_pairs = filtered_sic_df.dropna(subset=["gvkey"])
        company_to_gvkey = dict(
            zip(valid_gvkey_pairs["companyid"], valid_gvkey_pairs["gvkey"])
        )
        unique_whitelist_gvkeys = valid_gvkey_pairs["gvkey"].nunique()
        print(
            "GVKEY Mapping Dictionary Built:"
            f" {len(company_to_gvkey):,} links (Unique GVKEYs:"
            f" {unique_whitelist_gvkeys:,})"
        )
    else:
        print("GVKEY Column Not Found in SIC Mapping File.")

except Exception as e:
    print(f"ERROR in Stage 1: {e}")
    raise

# ==========================================
# Stage 2 : Reading & Merging Debt Files (USD 160 Filter)
# ==========================================
print("\n" + "=" * 80)
print("Stage 2 : Reading Debt Files")
print("=" * 80)

files = list(data_dir.glob("*.xlsx")) + list(data_dir.glob("*.csv"))
all_file_frames = []
file_summary = []

for file_path in files:
    if file_path.name in [
        output_analysis_excel.name,
        output_debug_excel.name,
        "Backtracked_Raw_Check_Report.csv",
        "Backtracked_Raw_Check_Report.xlsx",
    ]:
        continue

    print(f"Loading : {file_path.name}")
    try:
        if file_path.suffix.lower() == ".xlsx":
            df_raw = pd.read_excel(file_path)
        else:
            try:
                df_raw = pd.read_csv(file_path)
            except:
                df_raw = pd.read_csv(file_path, encoding="cp950")

        original_rows = len(df_raw)
        df_raw.columns = [extract_column_name(c) for c in df_raw.columns]

        required = {
            "companyid",
            "periodenddate",
            "componentid",
            "unittypeid",
            "dataitemvalue",
            "leveltypeid",
            "capitalstructuresubtypeid",
            "issuedcurrencyid",
            "descriptiontext",
        }
        missing = required - set(df_raw.columns)
        if len(missing) > 0:
            print(f"   [-] Missing Columns {missing}. Skip this file.")
            continue

        df_sub = df_raw[list(required)].copy()
        df_sub["companyid"] = df_sub["companyid"].apply(clean_to_string)

        # 僅保留美金 Issued Currency ID == 160
        df_sub["issuedcurrencyid"] = pd.to_numeric(
            df_sub["issuedcurrencyid"], errors="coerce"
        )
        before_currency = len(df_sub)
        df_sub = df_sub[df_sub["issuedcurrencyid"] == 160].copy()
        after_currency = len(df_sub)

        # 套用白名單過濾（剔除 SIC 為空值與非標產業公司）
        df_sub = df_sub[df_sub["companyid"].isin(valid_companies)].copy()
        after_whitelist = len(df_sub)

        # 正確精確對照 GVKEY（不強制用 companyid 補值，保持資料純淨）
        df_sub["gvkey"] = df_sub["companyid"].map(company_to_gvkey)
        df_sub["_Source_File"] = file_path.name

        if not df_sub.empty:
            all_file_frames.append(df_sub)

        file_summary.append({
            "Source_File": file_path.name,
            "Original_Rows": original_rows,
            "After_USD_160": after_currency,
            "After_SIC_Whitelist": after_whitelist,
            "Unique_Companies_In_File": df_sub["companyid"].nunique(),
        })
    except Exception as e:
        print(f"   [-] Error loading {file_path.name}: {e}")

if len(all_file_frames) == 0:
    raise Exception("Critical Error: No data frames loaded successfully.")

master_df = pd.concat(all_file_frames, ignore_index=True)

# 手動移除極端錯誤數據 903139141
if "componentid" in master_df.columns:
    master_df["comp_id_str"] = (
        master_df["componentid"]
        .astype(str)
        .str.strip()
        .str.split(".")
        .str[0]
    )
    master_df = master_df[master_df["comp_id_str"] != "903139141"].drop(
        columns=["comp_id_str"]
    ).copy()

master_df["Report_Year"] = master_df["periodenddate"].apply(parse_year)
master_df = master_df.dropna(subset=["Report_Year"])
master_df["Report_Year"] = master_df["Report_Year"].astype(int)

# 計算實質金額與債務分類
master_df["Actual_Amount"] = master_df.apply(
    lambda r: calc_actual_value(r["dataitemvalue"], r["unittypeid"]), axis=1
)
master_df["Debt_Category"] = master_df.apply(
    lambda r: map_debt_category(
        r["capitalstructuresubtypeid"], r["leveltypeid"]
    ),
    axis=1,
)

# 剔除無法分類/無效 (Debt_Category 為 None) 的觀測值
before_drop_unmapped = len(master_df)
master_df = master_df.dropna(subset=["Debt_Category"]).copy()
after_drop_unmapped = len(master_df)
print(
    "Removed Unmapped Debt Categories:"
    f" {before_drop_unmapped - after_drop_unmapped:,} rows"
)

# 計算總不重複公司數量
total_unique_gvkey = master_df["gvkey"].dropna().nunique()
total_unique_companyid = master_df["companyid"].nunique()

print(f"\nMaster Database Consolidated Successfully. Shape: {master_df.shape}")
print(f" Total Unique GVKEYs in Dataset      : {total_unique_gvkey:,}")
print(f" Total Unique CompanyIDs in Dataset  : {total_unique_companyid:,}")

# ==========================================
# Stage 3 : Generate Core Analysis Worksheets & Company Counts
# ==========================================
print("\n" + "=" * 80)
print("Stage 3 : Generating Core Worksheets & Company Counts")
print("=" * 80)

required_cats = ["CP", "DC", "TL", "CL", "SBN", "SUB", "Others"]

# Worksheets 1 & 2: Annual Totals
pivot_amount = master_df.pivot_table(
    index="Report_Year",
    columns="Debt_Category",
    values="Actual_Amount",
    aggfunc="sum",
).fillna(0)
for cat in required_cats:
    if cat not in pivot_amount.columns:
        pivot_amount[cat] = 0.0
pivot_amount = pivot_amount[required_cats]

pivot_percent = pivot_amount.div(pivot_amount.sum(axis=1), axis=0).fillna(0) * 100

final_amount_tbl = pivot_amount.T.round(0)
final_percent_tbl = pivot_percent.T.round(2)

# Worksheets 3 & 4: Firm-Year Panels (同時包含 companyid、gvkey 與 TotalDebt 欄位)
# 【關鍵修復】: dropna=True (預設)，防止記憶體崩潰 (Memory Error)
firm_year_pivot = master_df.pivot_table(
    index=["companyid", "gvkey", "Report_Year"],
    columns="Debt_Category",
    values="Actual_Amount",
    aggfunc="sum",
    dropna=True
).fillna(0)

for cat in required_cats:
    if cat not in firm_year_pivot.columns:
        firm_year_pivot[cat] = 0.0
firm_year_pivot = firm_year_pivot[required_cats]

# 剔除各類債務皆為零的觀測值
firm_year_pivot = firm_year_pivot[(firm_year_pivot != 0).any(axis=1)]

# 1. 建立金額面板 (Firm_Year_Amounts)
firm_year_amount_df = firm_year_pivot.copy()
firm_year_amount_df["TotalDebt"] = firm_year_amount_df.sum(axis=1)
firm_year_amount_df = (
    firm_year_amount_df.reset_index()
    .rename(columns={"Report_Year": "Year"})
    .round(0)
)

# 重新編排金額面板欄位順序：companyid -> gvkey -> Year -> 各類別 -> TotalDebt
amount_cols_order = (
    ["companyid", "gvkey", "Year"] + required_cats + ["TotalDebt"]
)
firm_year_amount_df = firm_year_amount_df[amount_cols_order]

# 2. 建立百分比面板 (Firm_Year_Percentages)
total_debt_series = firm_year_pivot.sum(axis=1)
firm_year_pct_pivot = firm_year_pivot.div(total_debt_series, axis=0).fillna(
    0
) * 100
firm_year_percent_df = (
    firm_year_pct_pivot.reset_index()
    .rename(columns={"Report_Year": "Year"})
    .round(2)
)

pct_cols_order = ["companyid", "gvkey", "Year"] + required_cats
firm_year_percent_df = firm_year_percent_df[pct_cols_order]

# 計算「每年不重複公司數」統計表
annual_companies = (
    master_df.groupby("Report_Year")
    .agg(
        Unique_GVKEY_Count=("gvkey", "nunique"),
        Unique_CompanyID_Count=("companyid", "nunique"),
        Total_Debt_Records=("Actual_Amount", "count"),
    )
    .reset_index()
    .rename(columns={"Report_Year": "Year"})
)

# 寫入主分析檔
try:
    with pd.ExcelWriter(output_analysis_excel, engine="openpyxl") as writer:
        final_amount_tbl.to_excel(writer, sheet_name="Actual_Amount_Data")
        final_percent_tbl.to_excel(writer, sheet_name="Percentage_Distribution")
        firm_year_amount_df.to_excel(
            writer, sheet_name="Firm_Year_Amounts", index=False
        )
        firm_year_percent_df.to_excel(
            writer, sheet_name="Firm_Year_Percentages", index=False
        )
        annual_companies.to_excel(
            writer, sheet_name="Unique_Companies_Summary", index=False
        )

        # 繪製原生 100% 堆疊圖
        ws_pct = writer.book["Percentage_Distribution"]
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.grouping = "percentStacked"
        chart.overlap = 100
        chart.title = (
            "Annual Debt Structure Percentage Distribution (SIC Filtered)"
        )
        chart.y_axis.title = "Percentage (100%)"
        chart.x_axis.title = "Year"

        data = Reference(
            ws_pct,
            min_col=2,
            min_row=1,
            max_col=ws_pct.max_column,
            max_row=ws_pct.max_row,
        )
        cats = Reference(ws_pct, min_col=1, min_row=2, max_row=ws_pct.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height, chart.width = 14, 22
        ws_pct.add_chart(chart, "A15")
    print(f" -> Main Analysis Report Saved: {output_analysis_excel.name}")
except PermissionError:
    print(" [!] Close the opening Excel file before running.")

# ==========================================
# Stage 4 : Advanced Debug & Auditing System
# ==========================================
print("\n" + "=" * 80)
print("Stage 4 : Building Comprehensive Debug Report")
print("=" * 80)

debug_df = (
    master_df[master_df["Report_Year"] == DEBUG_YEAR].copy()
    if DEBUG_YEAR is not None
    else master_df.copy()
)

# 1. Duplicate ComponentID
dup_mask = debug_df.duplicated(
    subset=["companyid", "Report_Year", "componentid"], keep=False
)
all_dups = debug_df[dup_mask].copy()

if not all_dups.empty:
    dup_summary = (
        all_dups.groupby(["companyid", "Report_Year", "_Source_File"])
        .size()
        .reset_index(name="Duplicate_Record_Count")
    )
    dup_output = dup_summary.sort_values(
        by="Duplicate_Record_Count", ascending=False
    ).head(500)
else:
    dup_output = pd.DataFrame([{
        "Status": (
            "Perfect! No Duplicate ComponentID found under same firm-year."
        )
    }])

# 2. Source File Analysis
source_file_df = (
    debug_df.groupby(["Report_Year", "_Source_File"])
    .agg(
        Record_Count=("Actual_Amount", "count"),
        Unique_CompanyIDs=("companyid", "nunique"),
    )
    .reset_index()
)


# 3. Description Text Analysis
def get_top_phrases(df, top_n=20):
    phrases = (
        df["descriptiontext"].dropna().astype(str).str.strip().str.upper()
    )
    short_phrases = phrases.apply(lambda x: " ".join(x.split()[:3]))
    return (
        short_phrases.value_counts()
        .head(top_n)
        .reset_index(name="Occurrences")
        .rename(columns={"index": "Phrase_Structure"})
    )


desc_audit_df = get_top_phrases(debug_df)

# 4. Top Debt Companies
annual_firm_total = (
    debug_df.groupby(["Report_Year", "companyid"])["Actual_Amount"]
    .sum()
    .reset_index(name="Firm_Annual_Total_Debt")
)
annual_market_total = (
    debug_df.groupby("Report_Year")["Actual_Amount"]
    .sum()
    .reset_index(name="Market_Annual_Total_Debt")
)
company_share_df = pd.merge(
    annual_firm_total, annual_market_total, on="Report_Year"
)
company_share_df["Market_Share_Percentage"] = (
    company_share_df["Firm_Annual_Total_Debt"]
    / company_share_df["Market_Annual_Total_Debt"]
) * 100
top_companies_df = company_share_df.sort_values(
    by=["Report_Year", "Market_Share_Percentage"], ascending=[True, False]
)

# 5. Top Debt Category Counts & Sums
cat_summary = (
    debug_df.groupby(["Report_Year", "Debt_Category"])
    .agg(
        Total_Amount=("Actual_Amount", "sum"),
        Record_Count=("Actual_Amount", "count"),
        Unique_Companies=("companyid", "nunique"),
    )
    .reset_index()
)

# 寫入 Debug 稽核報告檔
try:
    with pd.ExcelWriter(output_debug_excel, engine="openpyxl") as writer:
        pd.DataFrame(file_summary).to_excel(
            writer, sheet_name="Data_Ingestion_Summary", index=False
        )
        source_file_df.to_excel(
            writer, sheet_name="Source_File_Timeline", index=False
        )
        cat_summary.to_excel(
            writer, sheet_name="Category_Distribution", index=False
        )
        top_companies_df.to_excel(
            writer, sheet_name="Top_Debt_Companies", index=False
        )
        desc_audit_df.to_excel(
            writer, sheet_name="Description_String_Audit", index=False
        )
        dup_output.to_excel(
            writer, sheet_name="Duplicate_Component_Alert", index=False
        )

    print(
        " -> Diagnostic Debug Report Saved Successfully:"
        f" {output_debug_excel.name}"
    )
    print("=" * 80)
    print(" ALL PROCESSES CONCLUDED SUCCESSFULLY.")
except PermissionError:
    print(" [!] Close 'Debt_Debug_Report.xlsx' before re-running.")

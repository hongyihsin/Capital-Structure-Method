import pandas as pd
from pathlib import Path
import re
import numpy as np
import openpyxl
from openpyxl.chart import BarChart, Reference

# ==========================================
# 1. Settings (設定區)
# ==========================================
data_dir = Path(r"C:\ntutdata\debt")  
sic_excel_path = Path(r"C:\Users\olivi\Desktop\acer\ntut\1Seminar\data\GVKEY_debt structure_with SIC.xlsx")
output_analysis_excel = Path(r"C:\Users\olivi\Desktop\acer\ntut\1Seminar\data\Debt_Structure_Annual_Analysis.xlsx")

# ==========================================
# 2. Helper Functions (工具函數區)
# ==========================================
def extract_column_name(col):
    match = re.match(r"\((.*?)\)", str(col))
    return match.group(1).strip() if match else str(col).strip()

def clean_to_string(val):
    if pd.isna(val) or str(val).strip().lower() in ['nan', 'null', '']:
        return None
    return str(val).strip().split('.')[0]

def parse_year(date_val):
    """ 超級強化版年份提取：自動辨識各式各樣的台美日報表日期格式"""
    if pd.isna(date_val):
        return None
    date_str = str(date_val).strip()
    
    match = re.search(r"\b(19\d{2}|20\d{2})\b", date_str)
    if match:
        return int(match.group(1))
        
    try:
        return pd.to_datetime(date_str, errors='coerce').year
    except:
        return None

# ==========================================
# 3. Stage 1: Build Whitelist & GVKEY Mapping
# ==========================================
print("Stage 1: Processing SIC industry codes and building company whitelist...")
try:
    sic_df = pd.read_excel(sic_excel_path)
    sic_df.columns = [extract_column_name(c) for c in sic_df.columns]
    
    gvkey_col = None
    for c in sic_df.columns:
        if "companyid" in c.lower():
            sic_df = sic_df.rename(columns={c: "companyid"})
        if "sic" in c.lower() and c.lower() != "sic":
            sic_df = sic_df.rename(columns={c: "sic"})
        if "gvkey" in c.lower():
            gvkey_col = c

    sic_df["companyid"] = sic_df["companyid"].apply(clean_to_string)
    sic_df = sic_df.dropna(subset=["companyid"])
    
    if gvkey_col:
        sic_df = sic_df.rename(columns={gvkey_col: "gvkey"})
        sic_df["gvkey"] = sic_df["gvkey"].apply(clean_to_string)
        company_to_gvkey = dict(zip(sic_df["companyid"], sic_df["gvkey"]))
    else:
        print(" 警告：在對照檔中找不到 gvkey 欄位！將使用 companyid 代替。")
        company_to_gvkey = {}

    sic_df["sic"] = sic_df["sic"].fillna(-1)
    sic_df["sic_int"] = pd.to_numeric(sic_df["sic"], errors="coerce").fillna(-1).astype(int)
    
    filtered_sic_df = sic_df[
        ~((sic_df["sic_int"] >= 4900) & (sic_df["sic_int"] <= 4949)) & 
        ~((sic_df["sic_int"] >= 6000) & (sic_df["sic_int"] <= 6999))
    ].copy()
    
    valid_companies = set(filtered_sic_df["companyid"].dropna().unique())
    print(f"-> Total valid companies retained in whitelist: {len(valid_companies):,} companies.\n")
except Exception as e:
    print(f" Error loading SIC mapping file: {e}")
    raise e

# ==========================================
# 4. Stage 2: Merge All Files In-Memory
# ==========================================
print("Stage 2: Merging all excel/csv files and filtering by whitelist & Currency ID...")
files = list(data_dir.glob("*.xlsx")) + list(data_dir.glob("*.csv"))
all_file_frames = []

for file_path in files:
    if file_path.name in [output_analysis_excel.name, "Backtracked_Raw_Check_Report.csv", "Backtracked_Raw_Check_Report.xlsx"]:
        continue
        
    try:
        if file_path.suffix.lower() == ".xlsx":
            df_raw = pd.read_excel(file_path)
        else:
            try:
                df_raw = pd.read_csv(file_path, encoding='utf-8-sig')
            except:
                try:
                    df_raw = pd.read_csv(file_path, encoding='cp950')
                except:
                    df_raw = pd.read_csv(file_path, encoding='utf-8', errors='ignore')
                
        df_raw.columns = [str(c).strip() for c in df_raw.columns]
        df_raw.columns = [extract_column_name(c) for c in df_raw.columns]
        
        id_col = [c for c in df_raw.columns if "companyid" in c.lower() or "company_id" in c.lower() or c.lower() == "co_id"]
        date_col = [c for c in df_raw.columns if "periodenddate" in c.lower() or "period_end" in c.lower() or "report_date" in c.lower() or c.lower() == "date"]
        val_col = [c for c in df_raw.columns if "dataitemvalue" in c.lower() or "value" in c.lower() or "principal" in c.lower() or "amount" in c.lower()]
        unit_col = [c for c in df_raw.columns if "unittypeid" in c.lower() or "unit" in c.lower()]
        type_col = [c for c in df_raw.columns if "capitalstructuresubtypeid" in c.lower() or "subtype" in c.lower() or "debt_type" in c.lower()]
        level_col = [c for c in df_raw.columns if "leveltypeid" in c.lower() or "level" in c.lower() or "seniority" in c.lower()]
        currency_col = [c for c in df_raw.columns if "issuedcurrencyid" in c.lower() or "currency" in c.lower()]
        
        if not id_col or not date_col or not val_col:
            continue
            
        keep_cols = {id_col[0]: "companyid", date_col[0]: "periodenddate", val_col[0]: "dataitemvalue"}
        if unit_col: keep_cols[unit_col[0]] = "unittypeid"
        if type_col: keep_cols[type_col[0]] = "capitalstructuresubtypeid"
        if level_col: keep_cols[level_col[0]] = "leveltypeid"
        if currency_col: keep_cols[currency_col[0]] = "issuedcurrencyid"
        
        df_sub = df_raw[list(keep_cols.keys())].rename(columns=keep_cols).copy()
        df_sub["companyid"] = df_sub["companyid"].apply(clean_to_string)
        
        raw_len = len(df_sub)
        df_sub = df_sub[df_sub["companyid"].isin(valid_companies)].copy()
        whitelist_len = len(df_sub)
        
        if "issuedcurrencyid" in df_sub.columns:
            df_sub["issuedcurrencyid"] = pd.to_numeric(df_sub["issuedcurrencyid"], errors='coerce')
            df_sub = df_sub[df_sub["issuedcurrencyid"] == 160].copy()
            currency_len = len(df_sub)
        else:
            currency_len = whitelist_len
        
        if not df_sub.empty:
            if "unittypeid" not in df_sub.columns: df_sub["unittypeid"] = 0
            if "capitalstructuresubtypeid" not in df_sub.columns: df_sub["capitalstructuresubtypeid"] = 9
            if "leveltypeid" not in df_sub.columns: df_sub["leveltypeid"] = 1
            
            all_file_frames.append(df_sub)
            print(f"   + {file_path.name}: 原始 {raw_len:,} 筆 -> 白名單後 {whitelist_len:,} 筆 -> 幣別160後剩餘 {currency_len:,} 筆投遞成功")
                
    except Exception as e:
        print(f" 讀取 {file_path.name} 時發生嚴重錯誤: {e}")

# ==========================================
# 5. Stage 3: Data Processing & Mapping
# ==========================================
if all_file_frames:
    print("\nStage 3: Combining filtered data and performing diagnostic calculations...")
    master_df = pd.concat(all_file_frames, ignore_index=True)
    
    total_before_year = len(master_df)
    master_df["Report_Year"] = master_df["periodenddate"].apply(parse_year)
    
    parsed_years_found = master_df["Report_Year"].dropna().unique()
    print(f"   [INFO] All unique years successfully parsed in memory: {sorted(list(parsed_years_found))}")
    
    master_df = master_df.dropna(subset=["Report_Year"])
    total_after_year = len(master_df)
    print(f"   [DIAGNOSTIC] Records dropped due to date format issue (NaN): {total_before_year - total_after_year:,} records")
    
    master_df = master_df[master_df["Report_Year"] >= 1990].copy()
    
    master_df["gvkey"] = master_df["companyid"].map(company_to_gvkey)
    master_df["gvkey"] = master_df["gvkey"].fillna(master_df["companyid"])
    
    def calc_actual_value(row):
        val = float(row["dataitemvalue"]) if pd.notna(row["dataitemvalue"]) else 0.0
        unit = row["unittypeid"]
        if unit == 1: return val * 1000
        elif unit == 2: return val * 1000000
        return val
    master_df["Actual_Amount"] = master_df.apply(calc_actual_value, axis=1)
    
    def map_debt_category(row):
        tid = row["capitalstructuresubtypeid"]
        lid = row["leveltypeid"]
        if tid == 1: return "CP"
        elif tid == 2: return "DC"
        elif tid == 3: return "TL"
        elif tid == 5: return "CL"
        elif tid in [6, 7, 9]: return "other"
        elif tid == 4:
            if lid == 1: return "SBN"
            elif lid in [2, 3, 4, 5, 6, 7]: return "sub"
        return "other"
    master_df["Debt_Category"] = master_df.apply(map_debt_category, axis=1)
    
    # ==========================================
    # 6. Stage 4: Generate Worksheets Data (包含全零過濾)
    # ==========================================
    print("Stage 4: Preparing data for 4 different worksheets...")
    pivot_annual_amount = master_df.pivot_table(
        index="Report_Year", columns="Debt_Category", values="Actual_Amount", aggfunc="sum"
    ).fillna(0)
    
    required_cats = ["CP", "DC", "TL", "CL", "other", "SBN", "sub"]
    for cat in required_cats:
        if cat not in pivot_annual_amount.columns:
            pivot_annual_amount[cat] = 0.0
    pivot_annual_amount = pivot_annual_amount[required_cats]
    
    pivot_annual_percent = pivot_annual_amount.div(pivot_annual_amount.sum(axis=1), axis=0).fillna(0) * 100
    final_amount_tbl = pivot_annual_amount.T.round(0)
    final_percent_tbl = pivot_annual_percent.T.round(2)
    
    # 建立公司年度樞紐表
    firm_year_pivot = master_df.pivot_table(
        index=["gvkey", "Report_Year"], columns="Debt_Category", values="Actual_Amount", aggfunc="sum"
    ).fillna(0)
    for cat in required_cats:
        if cat not in firm_year_pivot.columns:
            firm_year_pivot[cat] = 0.0
    firm_year_pivot = firm_year_pivot[required_cats]
    
    # 💡【核心修改】只要任一債務類別不為 0 才保留，徹底剔除全零列
    firm_year_pivot = firm_year_pivot[(firm_year_pivot != 0).any(axis=1)]
    
    # 轉換成最終輸出的 DataFrames
    firm_year_amount_df = firm_year_pivot.reset_index().rename(columns={"Report_Year": "year"}).round(0)
    total_debt_series = firm_year_pivot.sum(axis=1)
    firm_year_pct_pivot = firm_year_pivot.div(total_debt_series, axis=0).fillna(0) * 100
    firm_year_percent_df = firm_year_pct_pivot.reset_index().rename(columns={"Report_Year": "year"}).round(2)

    # ==========================================
    # 7. Stage 5: Export to Excel with Native Charts
    # ==========================================
    print(f"\nStage 5: Writing 4 sheets into Excel...")
    try:
        with pd.ExcelWriter(output_analysis_excel, engine='openpyxl') as writer:
            final_amount_tbl.to_excel(writer, sheet_name="Actual_Amount_Data")
            final_percent_tbl.to_excel(writer, sheet_name="Percentage_Distribution")
            firm_year_amount_df.to_excel(writer, sheet_name="Firm_Year_Amounts", index=False)
            firm_year_percent_df.to_excel(writer, sheet_name="Firm_Year_Percentages", index=False)
            
            workbook = writer.book
            ws_pct = workbook["Percentage_Distribution"]
            
            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.grouping = "percentStacked"
            chart.overlap = 100
            chart.title = "Annual Debt Structure Percentage Distribution (SIC & Currency Filtered, 1990+)"
            chart.y_axis.title = "Percentage (100%)"
            chart.x_axis.title = "Debt Category"
            
            max_col = ws_pct.max_column
            max_row = ws_pct.max_row
            data = Reference(ws_pct, min_col=2, min_row=1, max_col=max_col, max_row=max_row)
            cats = Reference(ws_pct, min_col=1, min_row=2, max_row=max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.height = 14
            chart.width = 22
            ws_pct.add_chart(chart, "A11")
        print("\n ALL 4 WORKSHEETS GENERATED SUCCESSFULLY IN EXCEL!")
    except PermissionError:
        print("\n 錯誤：請先關閉開著的 Excel 檔案再重新執行程式！")
else:
    print("\n 處理失敗。")